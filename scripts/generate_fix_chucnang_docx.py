import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles & Colors
    # Primary: #0F3A66 (Deep Navy)
    # Secondary / Teal: #0284C7 (Medical Teal/Cyan)
    # Slate Text: #1E293B
    # Light Slate: #64748B
    # Danger / Red: #DC2626
    # Warning / Amber: #D97706
    # Success / Green: #16A34A
    
    # Helper Functions
    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

    def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_cell_borders(cell, top='single', bottom='single', left='none', right='none', color='CBD5E1', sz='4'):
        tcPr = cell._tc.get_or_add_tcPr()
        borders_xml = f'<w:tcBorders {nsdecls("w")}>'
        borders_xml += f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        borders_xml += f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        borders_xml += f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        borders_xml += f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        borders_xml += '</w:tcBorders>'
        tcPr.append(parse_xml(borders_xml))

    def add_doc_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 58, 102) # #0F3A66
        return p

    def add_doc_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(18)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(2, 132, 199) # #0284C7
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 58, 102)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(2, 132, 199)
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_p(text, bold_prefix=None, italic=False, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.name = 'Arial'
            run_b.font.size = Pt(10.5)
            run_b.font.bold = True
            run_b.font.color.rgb = RGBColor(30, 41, 59)
        if text:
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(10.5)
            run.font.italic = italic
            run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_bullet(text, bold_prefix=None, level=0):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.name = 'Arial'
            run_b.font.size = Pt(10)
            run_b.font.bold = True
            run_b.font.color.rgb = RGBColor(30, 41, 59)
        if text:
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_callout(title, text, callout_type="info"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)

        bg_color = "F0F9FF"  # Default light blue
        border_color = "0284C7"
        title_color = RGBColor(2, 132, 199)

        if callout_type == "danger":
            bg_color = "FEF2F2"
            border_color = "DC2626"
            title_color = RGBColor(220, 38, 38)
        elif callout_type == "warning":
            bg_color = "FFFBEB"
            border_color = "D97706"
            title_color = RGBColor(217, 119, 6)
        elif callout_type == "success":
            bg_color = "F0FDF4"
            border_color = "16A34A"
            title_color = RGBColor(22, 163, 74)

        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        # Set thick left border, no others
        tcPr = cell._tc.get_or_add_tcPr()
        borders_xml = f'<w:tcBorders {nsdecls("w")}>'
        borders_xml += f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        borders_xml += '<w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
        borders_xml += '</w:tcBorders>'
        tcPr.append(parse_xml(borders_xml))

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        run_t = p.add_run(f"📌 {title}\n" if title else "")
        run_t.font.name = 'Arial'
        run_t.font.size = Pt(10.5)
        run_t.font.bold = True
        run_t.font.color.rgb = title_color

        run_b = p.add_run(text)
        run_b.font.name = 'Arial'
        run_b.font.size = Pt(10)
        run_b.font.color.rgb = RGBColor(30, 41, 59)

        # Space after table
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(4)

    def add_styled_table(headers, rows_data, col_widths=None):
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Header Row
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            cell = hdr_cells[i]
            cell.text = header_text
            set_cell_background(cell, "0F3A66")
            set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
            set_cell_borders(cell, top='single', bottom='single', left='single', right='single', color='CBD5E1', sz='4')
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.name = 'Arial'
            p.runs[0].font.size = Pt(10)
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, row_values in enumerate(rows_data):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_values):
                cell = row_cells[c_idx]
                cell.text = str(val)
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                set_cell_borders(cell, top='single', bottom='single', left='single', right='single', color='E2E8F0', sz='4')
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                if len(p.runs) > 0:
                    p.runs[0].font.name = 'Arial'
                    p.runs[0].font.size = Pt(9.5)
                    p.runs[0].font.color.rgb = RGBColor(51, 65, 85)

        # Apply Column Widths if provided
        if col_widths:
            for row in table.rows:
                for idx, width in enumerate(col_widths):
                    row.cells[idx].width = Inches(width)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(6)
        return table

    # -------------------------------------------------------------
    # DOCUMENT CONTENT GENERATION
    # -------------------------------------------------------------
    
    # Title & Header
    add_doc_title("TÀI LIỆU ĐẶC TẢ CHI TIẾT YÊU CẦU NÂNG CẤP & HOÀN THIỆN CHỨC NĂNG")
    add_doc_subtitle("HỆ THỐNG ALLERCARE AI V2 - PHIÊN BẢN CẢI TIẾN NGHIỆP VỤ 10.3")

    # Document Metadata Box
    meta_headers = ["Thuộc tính", "Thông tin chi tiết"]
    meta_rows = [
        ["Hệ thống", "AllerCare AI - Nền tảng theo dõi từ xa & Giảm thiểu sai sót sử dụng thuốc"],
        ["Chuyên khoa áp dụng", "Da liễu - Dị ứng Miễn dịch Lâm sàng"],
        ["Mục đích tài liệu", "Chuẩn hóa toàn diện các yêu cầu chức năng, luồng tương tác UX, mô hình dữ liệu và API cho bản cập nhật 10.3 theo yêu cầu chuyên môn y tế."],
        ["Đối tượng áp dụng", "Nhóm phát triển Backend (.NET/FastAPI), Frontend (Next.js/React), AI Engineer & Chuyên gia Y tế"],
        ["Trạng thái tài liệu", "Chính thức ban hành (Approved for Development & Testing)"]
    ]
    add_styled_table(meta_headers, meta_rows, col_widths=[2.0, 4.5])

    # -------------------------------------------------------------
    # SECTION 1: TỔNG QUAN VÀ BỐI CẢNH NÂNG CẤP
    # -------------------------------------------------------------
    add_heading_1("1. TỔNG QUAN VÀ MỤC TIÊU NÂNG CẤP (PHIÊN BẢN 10.3)")
    add_p("Trong quá trình vận hành thực tế và đánh giá lâm sàng hệ thống AllerCare AI V2, nhóm chuyên môn y tế và người dùng đã phát hiện một số điểm nghẽn (pain points) liên quan đến trải nghiệm tuân thủ của Người bệnh và khả năng giám sát, kê đơn an toàn của Bác sĩ. Tài liệu này chuẩn hóa toàn diện 5 nhóm yêu cầu cải tiến trọng yếu:")
    
    add_bullet("Hệ thống nhắc nhở đa kênh (Push/Alarm/Widget màn hình chính) kèm trạng thái cảnh báo trực quan (viền nháy đỏ) và cơ chế check-in 1 chạm (Quick chips) gửi dữ liệu tức thời sang Bác sĩ.", "1. Phân hệ Bệnh nhân - Tuân thủ điều trị: ")
    add_bullet("Bổ sung rõ ràng nút đóng đợt điều trị (Lưu trữ lịch sử) và nút xóa chẩn đoán (Xóa nhầm có ghi vết Audit log an toàn).", "2. Phân hệ Bác sĩ - Quản lý bệnh lý: ")
    add_bullet("Hiển thị Badge số đếm đỏ đại diện cho các thay đổi mới (Triệu chứng mới, thuốc OTC tự dùng, nhật ký check-in) ngay phía sau tên từng bệnh nhân trên Sidebar.", "3. Phân hệ Bác sĩ - Trung tâm thông báo: ")
    add_bullet("Tách bạch danh mục thuốc điều trị theo từng mặt bệnh riêng biệt (mỗi bệnh có danh mục thuốc riêng), đồng thời duy trì danh mục tổng thể ở trang chủ.", "4. Phân hệ Bác sĩ - Phân loại thuốc theo bệnh: ")
    add_bullet("Autocomplete tìm kiếm thuốc ngay từ ký tự đầu tiên kết hợp kiểm tra an toàn AI tại chỗ (Tương tác thuốc - thuốc, tương tác thuốc - bệnh lý/dị ứng) trước khi lưu đơn.", "5. Phân hệ Bác sĩ - Kê thuốc thông minh & AI Guardrails: ")

    add_callout(
        "Ý Nghĩa Lâm Sàng & An Toàn Y Khoa",
        "Các cải tiến này đảm bảo nguyên tắc cốt lõi của AllerCare AI: Giảm thiểu sai sót dùng thuốc, nâng cao tỷ lệ tuân thủ điều trị ngoại trú của bệnh nhân da liễu - dị ứng, và cung cấp công cụ trợ lý AI an toàn, giải thích rõ ràng cho Bác sĩ điều trị.",
        "info"
    )

    # -------------------------------------------------------------
    # SECTION 2: PHÂN HỆ NGƯỜI BỆNH (PATIENT PORTAL / PWA)
    # -------------------------------------------------------------
    add_heading_1("2. ĐẶC TẢ CHI TIẾT PHÂN HỆ NGƯỜI BỆNH (PATIENT APP)")
    
    add_heading_2("2.1. Chức năng Nhắc nhở Uống thuốc & Báo thức Thông minh")
    add_p("Quy tắc khởi đầu ngày mới của bệnh nhân da liễu/dị ứng: Uống thuốc đúng giờ và cập nhật tình trạng bệnh hiện tại là hai hành động quan trọng nhất cần được hoàn thành đầu tiên trong ngày.")
    
    add_heading_3("a) Cơ chế Nhắc nhở & Báo thức (Alarm / Push Notification):")
    add_bullet("Hệ thống tự động kích hoạt thông báo đẩy (Web Push / Mobile App Notification) và chuông báo thức theo đúng khung giờ quy định trong đơn thuốc của Bác sĩ (ví dụ: 07:30 Sáng, 12:30 Trưa, 20:00 Tối).", "• Cơ chế báo thức: ")
    add_bullet("Thông báo hiển thị dạng Modal Full-screen hoặc Sticky Banner cố định trên màn hình chính của ứng dụng với âm thanh/rung nhẹ nhàng.", "• Trải nghiệm tương tác: ")
    add_bullet("Bệnh nhân có thể bấm 'Đã uống', 'Nhắc lại sau 10 phút (Snooze)' hoặc 'Chưa uống / Bỏ qua' (kèm lý do: Quên, Bận, Buồn nôn, Hết thuốc).", "• Nút bấm thao tác nhanh: ")

    add_heading_3("b) Widget Màn hình chính & Cảnh báo Nháy đỏ (Pulsing Red Alert):")
    add_bullet("Nếu đến giờ uống thuốc hoặc chưa hoàn thành khai báo đầu ngày, Thẻ trạng thái trên Dashboard sẽ kích hoạt hiệu ứng Nháy đỏ (Pulsing Red Border) kèm dòng chữ: '⚠️ BẠN CẦN UỐNG THUỐC VÀ KHAI BÁO SỨC KHỎE HÔM NAY'.", "• Trạng thái Chưa hoàn thành: ")
    add_bullet("Khi bệnh nhân đã xác nhận uống thuốc và hoàn thành check-in, thẻ trạng thái chuyển ngay sang màu Xanh lục (Success Green) hiển thị: '✅ Đã hoàn thành nhiệm vụ sức khỏe ngày hôm nay'.", "• Trạng thái Đã hoàn thành: ")

    add_heading_2("2.2. Chức năng Check-in Tình trạng Bệnh Hằng ngày (Daily Health Check-in)")
    add_p("Thay vì bắt người bệnh phải nhập liệu phức tạp, giao diện cung cấp cơ chế cập nhật siêu tốc với 2 phương thức song song:")
    
    checkin_headers = ["Phương thức", "Hình thức giao diện", "Mô tả chi tiết & Dữ liệu thu thập"]
    checkin_rows = [
        [
            "1. Quick-Select Chips\n(Một chạm siêu tốc)",
            "Hàng nút bấm (Pill Chips) phân màu trực quan",
            "Bệnh nhân chỉ cần chạm nhanh vào các trạng thái gợi ý sẵn:\n• 🟢 Ổn định / Đỡ ngứa / Giảm ban đỏ\n• 🟡 Còn ngứa nhẹ / Da khô rát\n• 🔴 Ngứa dữ dội / Mẩn đỏ lan rộng\n• ⚠️ Có mụn nước / Chảy dịch / Phù nề\n• 🤢 Tác dụng phụ: Buồn nôn, Chóng mặt, Mệt mỏi"
        ],
        [
            "2. Nhập chi tiết\n(Text & Hình ảnh)",
            "Ô nhập văn bản tự do + Nút chụp/tải ảnh da",
            "Dành cho trường hợp bệnh nhân muốn mô tả chi tiết diễn biến bất thường hoặc tải ảnh chụp vùng da tổn thương mới để Bác sĩ xem xét trực quan."
        ]
    ]
    add_styled_table(checkin_headers, checkin_rows, col_widths=[1.5, 2.0, 3.0])

    add_heading_2("2.3. Đường ống Đồng bộ Tức thời sang Bác sĩ (Real-time Pipeline)")
    add_p("Ngay sau khi Bệnh nhân bấm [GỬI BÁO CÁO NGÀY], hệ thống thực hiện:")
    add_bullet("Ghi nhận nhật ký tuân thủ (Compliance Log) vào cơ sở dữ liệu.", "1. Lưu trữ: ")
    add_bullet("AI chạy phân loại Triage (Xanh: Ổn định, Vàng: Cần theo dõi, Đỏ: Cảnh báo khẩn cấp).", "2. Đánh giá AI: ")
    add_bullet("Đẩy thông báo tức thời (WebSocket / Push Event) sang Dashboard của Bác sĩ phụ trách, đồng thời tăng Badge đếm thông báo chưa xem trên hồ sơ bệnh nhân đó.", "3. Bắn tín hiệu sang Bác sĩ: ")

    # -------------------------------------------------------------
    # SECTION 3: PHÂN HỆ BÁC SĨ (DOCTOR CLINICAL DASHBOARD)
    # -------------------------------------------------------------
    add_heading_1("3. ĐẶC TẢ CHI TIẾT PHÂN HỆ BÁC SĨ (DOCTOR PORTAL)")

    add_heading_2("3.1. Chức năng Quản lý Bệnh lý: Nút Kết thúc Điều trị & Nút Xóa Bệnh")
    add_p("Để đảm bảo tính chính xác và an toàn y khoa, hệ thống phân định rõ ràng 2 nghiệp vụ:")

    manage_headers = ["Nghiệp vụ", "Nút bấm / Thao tác", "Mục đích lâm sàng", "Hành vi hệ thống & Dữ liệu"]
    manage_rows = [
        [
            "Kết thúc điều trị\n(Mark as Resolved)",
            "Nút [Kết thúc đợt điều trị]\n(Màu Xanh dương/Xanh lá)",
            "Bệnh nhân đã khỏi bệnh hoặc kết thúc phác đồ điều trị cho mặt bệnh cụ thể.",
            "• Chuyển trạng thái bệnh sang 'Đã kết thúc / Đã khỏi'.\n• Thuốc thuộc bệnh này chuyển sang trạng thái 'Đã ngừng / Hoàn thành'.\n• Hồ sơ bệnh sử VẪN ĐƯỢC LƯU VĨNH VIỄN để phục vụ tra cứu tương tác thuốc và bệnh sử sau này."
        ],
        [
            "Xóa bệnh điều trị\n(Void / Error Deletion)",
            "Nút [Xóa mặt bệnh]\n(Màu Đỏ cảnh báo)",
            "Bác sĩ hoặc nhân viên y tế nhập nhầm chẩn đoán / ghi nhận sai thông tin ban đầu.",
            "• Hiển thị Modal cảnh báo: Bắt buộc Bác sĩ nhập lý do xóa.\n• Thực hiện Soft-delete (ẩn khỏi giao diện điều trị hiện tại).\n• Ghi lại Audit Log (Ai xóa, lúc nào, lý do gì) phục vụ thanh tra y khoa."
        ]
    ]
    add_styled_table(manage_headers, manage_rows, col_widths=[1.4, 1.6, 1.7, 1.8])

    add_heading_2("3.2. Trung tâm Thông báo & Badge Đếm Cập nhật Mới Thời Gian Thực")
    add_p("Giúp Bác sĩ nhận diện ngay lập tức những bệnh nhân đang có biến động sức khỏe mà không cần phải mở từng hồ sơ ra kiểm tra thủ công.")

    add_heading_3("a) Hiển thị Badge số đếm tại Danh sách Bệnh nhân (Sidebar bên trái):")
    add_bullet("Phía sau tên mỗi bệnh nhân, hiển thị Badge hình tròn màu đỏ kèm số lượng cập nhật mới chưa đọc (ví dụ: 'Nguyễn Văn A [2]').", "• Vị trí & Định dạng: ")
    add_bullet("Tổng hợp từ 4 nguồn sự kiện: (1) Nhật ký check-in hàng ngày mới; (2) Triệu chứng mới khai báo; (3) Thuốc mới tự mua (OTC) người bệnh tự khai báo; (4) Báo cáo tác dụng phụ sau dùng thuốc.", "• Công thức tính số Badge: ")
    add_bullet("Nếu có sự kiện phân luồng ĐỎ (Khẩn cấp/Dị ứng nặng), Badge sẽ chuyển sang hiệu ứng Nhấp nháy đỏ kèm icon cảnh báo nguy hiểm ⚠️.", "• Điểm nhấn khẩn cấp: ")

    add_heading_3("b) Xem chi tiết và Reset Badge:")
    add_p("Khi Bác sĩ nhấp vào tên bệnh nhân, một ngăn kéo (Drawer/Modal) 'NHẬT KÝ BIẾN ĐỘNG MỚI' hiển thị danh sách các thay đổi theo dòng thời gian. Sau khi Bác sĩ xem và bấm [ĐÃ XEM], Badge số đếm sẽ tự động giảm về 0.")

    add_heading_2("3.3. Phân tách Danh mục Thuốc theo Từng Loại Bệnh Điều trị")
    add_p("Khắc phục triệt để nhược điểm hiển thị gộp chung thuốc, giúp Bác sĩ dễ dàng quản lý phác đồ đa bệnh lý (ví dụ: Bệnh nhân vừa điều trị Viêm da cơ địa, vừa điều trị Mề đay dị ứng).")

    add_callout(
        "Cấu Trúc Hiển Thị Thuốc Chuẩn Hóa",
        "1. Trang chủ / Tổng quan Hồ sơ: Duy trì widget 'Danh mục thuốc đang điều trị (Tổng thể)' - Cho cái nhìn toàn diện về tất cả các thuốc bệnh nhân đang nạp vào cơ thể.\n2. Trong từng Tab / Thẻ Bệnh lý riêng: Bổ sung tiểu mục 'Danh mục thuốc điều trị - [Tên Bệnh / STT Bệnh]' - Chỉ hiển thị danh sách các thuốc được kê riêng cho mặt bệnh đó kèm liều lượng và thời gian dùng.",
        "success"
    )

    add_heading_2("3.4. Kê Thuốc Thông Minh (Autocomplete) & Cảnh Báo An Toàn AI Thời Gian Thực")
    add_p("Hỗ trợ Bác sĩ ra quyết định lâm sàng nhanh chóng, chính xác và ngăn ngừa sai sót thuốc ngay từ thời điểm gõ đơn.")

    add_heading_3("a) Tính năng Tìm kiếm Autocomplete Siêu tốc:")
    add_bullet("Chỉ cần gõ 1 ký tự đầu tiên (ví dụ gõ 'c' -> gợi ý Cetirizine, Clobetasol, Cefuroxime, Clarithromycin...).", "• Tốc độ phản hồi: ")
    add_bullet("Mỗi mục trong danh sách xổ xuống (Dropdown) hiển thị rõ: Tên biệt dược, Hoạt chất gốc (Generic name), Hàm lượng (5mg, 10mg, 500mg...), Dạng bào chế (Viên nén, Kem bôi, Dung dịch) và Nhóm dược lý.", "• Thông tin phong phú: ")

    add_heading_3("b) Cảnh báo An toàn AI Thời gian thực (Real-time AI Guardrails):")
    add_p("Ngay khi Bác sĩ chọn 1 loại thuốc từ danh sách gợi ý, AI Engine sẽ thực hiện kiểm tra ngầm lập tức và đưa ra cảnh báo đa tầng trước khi bác sĩ lưu đơn:")

    safety_headers = ["Tầng kiểm tra", "Đối tượng so khớp", "Nội dung cảnh báo AI & Hành động khuyến nghị"]
    safety_rows = [
        [
            "1. Tương tác Thuốc - Thuốc\n(Drug - Drug Interaction)",
            "Thuốc dự kiến kê VS Toàn bộ thuốc đang dùng (kể cả thuốc ở bệnh khác & thuốc OTC)",
            "• Cảnh báo tương tác chéo, hiệp đồng độc tính hoặc giảm hấp thu.\n• Hiển thị: Mức độ nghiêm trọng (Chống chỉ định / Nghiêm trọng / Trung bình), Cơ chế tương tác và Khuyến nghị đổi thuốc hoặc giãn cách giờ uống."
        ],
        [
            "2. Tương tác Thuốc - Bệnh lý\n(Drug - Disease Contraindication)",
            "Thuốc dự kiến kê VS Tiền sử bệnh nền (Suy gan, Suy thận, Viêm loét dạ dày, Hen phế quản, Thai kỳ)",
            "• Cảnh báo nếu hoạt chất chuyển hóa qua gan/thận hoặc có chống chỉ định với bệnh nền hiện tại của người bệnh."
        ],
        [
            "3. Tương tác Thuốc - Dị ứng\n(Drug - Allergy History)",
            "Thuốc dự kiến kê VS Tiền sử dị ứng thuốc đã ghi nhận trong hồ sơ",
            "• Cảnh báo dị ứng chéo (Cross-reactivity) cùng nhóm (ví dụ: Dị ứng Penicillin cảnh báo khi kê Cephalosporin thế hệ 1)."
        ]
    ]
    add_styled_table(safety_headers, safety_rows, col_widths=[1.8, 2.0, 2.7])

    # -------------------------------------------------------------
    # SECTION 4: THIẾT KẾ KỸ THUẬT & API CONTRACTS
    # -------------------------------------------------------------
    add_heading_1("4. THIẾT KẾ DỮ LIỆU & API KỸ THUẬT")
    
    add_heading_2("4.1. Cập nhật Mô hình Dữ liệu (Data Model / Schema)")
    add_p("Bổ sung và chuẩn hóa các trường dữ liệu mới trong cơ sở dữ liệu:")
    add_bullet("Bổ sung trường status ('active', 'resolved', 'deleted'), resolved_at (timestamp), deleted_reason (text), deleted_by (UUID).", "• Bảng MedicalConditions (Bệnh lý): ")
    add_bullet("Bổ sung khóa ngoại condition_id để ánh xạ chính xác thuốc này được kê cho bệnh lý nào.", "• Bảng Prescriptions / MedicationItems (Thuốc): ")
    add_bullet("Lưu trữ compliance_status ('taken', 'skipped', 'snoozed'), quick_chips (JSON array), notes (text), image_url (text), triage_color ('green', 'yellow', 'red').", "• Bảng DailyCheckins (Nhật ký sức khỏe ngày): ")
    add_bullet("Lưu trữ doctor_id, patient_id, unread_count (int), last_event_type ('checkin', 'symptom', 'otc_drug'), last_event_time (timestamp).", "• Bảng DoctorPatientAlertSummary: ")

    add_heading_2("4.2. Danh sách API Endpoints Cốt lõi")
    api_headers = ["Phương thức & Endpoint", "Mô tả chức năng", "Dữ liệu Request / Response chính"]
    api_rows = [
        [
            "POST /api/patient/check-in",
            "Bệnh nhân gửi khai báo đầu ngày",
            "Req: { taken_meds: bool, chips: ['do_ngua'], notes: '...', image: base64 }\nRes: { success: true, triage: 'green', synced_to_doctor: true }"
        ],
        [
            "POST /api/doctor/conditions/{id}/resolve",
            "Bác sĩ kết thúc đợt điều trị bệnh",
            "Req: { condition_id: 'uuid', notes: 'Khỏi bệnh hoàn toàn' }\nRes: { status: 'resolved', archived_at: '2026-10-03T...' }"
        ],
        [
            "DELETE /api/doctor/conditions/{id}",
            "Bác sĩ xóa bệnh do ghi nhầm",
            "Req: { reason: 'Nhập nhầm từ hồ sơ bệnh nhân khác' }\nRes: { success: true, audit_logged: true }"
        ],
        [
            "GET /api/doctor/patient-updates",
            "Lấy danh sách Badge số đếm chưa đọc",
            "Res: [{ patient_id: 'p1', unread_events: 3, has_emergency: false }, ...]"
        ],
        [
            "GET /api/drugs/autocomplete?q={keyword}",
            "Tìm kiếm gợi ý thuốc tức thì",
            "Res: [{ id: 1, brand_name: 'Cetirizin 10mg', active_substance: 'Cetirizine', form: 'Viên' }]"
        ],
        [
            "POST /api/ai/prescribing-safety-check",
            "Kiểm tra an toàn AI tại chỗ khi kê thuốc",
            "Req: { patient_id: 'p1', target_drug_id: 102, condition_id: 'c1' }\nRes: { severity: 'warning', interactions: [...], disease_warnings: [...] }"
        ]
    ]
    add_styled_table(api_headers, api_rows, col_widths=[2.3, 1.7, 2.5])

    # -------------------------------------------------------------
    # SECTION 5: KỊCH BẢN KIỂM THỬ & TIÊU CHÍ NGHIỆM THU
    # -------------------------------------------------------------
    add_heading_1("5. KỊCH BẢN KIỂM THỬ VÀ TIÊU CHÍ NGHIỆM THU (ACCEPTANCE CRITERIA)")
    
    test_headers = ["STT", "Chức năng kiểm thử", "Hành động thực hiện (Test Steps)", "Kết quả mong đợi (Expected Outcome)"]
    test_rows = [
        [
            "TC-01",
            "Nhắc uống thuốc & Báo thức",
            "Đến giờ uống thuốc theo đơn, mở ứng dụng trên điện thoại.",
            "Hiện Modal thông báo + chuông. Widget màn hình chính nháy viền đỏ nếu chưa bấm 'Đã uống'."
        ],
        [
            "TC-02",
            "Check-in 1 chạm (Quick chips)",
            "Chọn chip 'Đỡ ngứa' + 'Không có triệu chứng phụ' -> Bấm Gửi.",
            "Trạng thái Dashboard chuyển Xanh. Dữ liệu đồng bộ sang Bác sĩ trong vòng < 1 giây."
        ],
        [
            "TC-03",
            "Badge đếm số sau tên bệnh nhân",
            "Bệnh nhân A gửi check-in và khai báo thuốc mới tự mua.",
            "Sidebar Bác sĩ hiển thị badge số [2] màu đỏ ngay sau tên Bệnh nhân A."
        ],
        [
            "TC-04",
            "Kết thúc vs Xóa bệnh lý",
            "1. Bấm Kết thúc điều trị -> Kiểm tra trạng thái Lưu trữ.\n2. Bấm Xóa bệnh -> Kiểm tra hộp thoại nhập lý do và Audit log.",
            "Bệnh kết thúc vẫn giữ lịch sử tra cứu. Bệnh xóa bị ẩn khỏi UI và ghi nhận log an toàn."
        ],
        [
            "TC-05",
            "Danh mục thuốc theo từng bệnh",
            "Mở Tab 'Viêm da cơ địa' và Tab 'Mày đay mạn'.",
            "Thuốc của bệnh nào hiển thị đúng trong Tab bệnh đó; Trang chủ hiển thị tổng hợp toàn bộ thuốc."
        ],
        [
            "TC-06",
            "Autocomplete & Cảnh báo AI",
            "Gõ chữ 'c' tại ô kê thuốc -> Chọn Ciprofloxacin cho bệnh nhân đang dùng Theophylline.",
            "Dropdown gợi ý ngay từ ký tự đầu. AI hiển thị cảnh báo ĐỎ tương tác thuốc mức độ nặng kèm lý giải."
        ]
    ]
    add_styled_table(test_headers, test_rows, col_widths=[0.6, 1.8, 2.1, 2.0])

    add_heading_2("5.1. Định nghĩa Hoàn thành (Definition of Done - DoD)")
    add_bullet("Tất cả 5 nhóm chức năng được hiện thực hóa đầy đủ trên cả Backend (.NET/FastAPI) và Frontend (Next.js/React).", "1. Tính đầy đủ: ")
    add_bullet("AI Guardrails phản hồi kiểm tra tương tác trong thời gian dưới 500ms, không gây độ trễ khi Bác sĩ thao tác.", "2. Hiệu năng & Tốc độ: ")
    add_bullet("Tuyệt đối không để AI tự ý thay đổi đơn thuốc hoặc chẩn đoán; Bác sĩ là người duy nhất ra quyết định cuối cùng.", "3. Ranh giới an toàn y tế: ")
    add_bullet("Vượt qua toàn bộ 6 kịch bản kiểm thử lâm sàng (TC-01 đến TC-06) không có lỗi nghiêm trọng.", "4. Kiểm thử nghiệm thu: ")

    # Save Document
    target_path = "/Users/macbook/AllerCare_AI_V2/src/AppHost/obj/FIX_CHUCNANG_10.3.docx"
    doc.save(target_path)
    print(f"Document successfully created and saved to: {target_path}")

if __name__ == "__main__":
    create_document()
